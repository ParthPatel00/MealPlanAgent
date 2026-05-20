#!/usr/bin/env python3
"""Generate a minimalist 12-slide presentation for CMPE 258 project."""

import os
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = Path(__file__).resolve().parent.parent
os.chdir(ROOT)
CHARTS = Path("report/charts")
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

BG = RGBColor(0x1A, 0x1A, 0x2E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0x00, 0xD2, 0xFF)
LIGHT_GRAY = RGBColor(0xBB, 0xBB, 0xBB)
DARK_CARD = RGBColor(0x25, 0x25, 0x3E)


def set_bg(slide):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = BG


def add_text(slide, left, top, width, height, text, size=18, color=WHITE,
             bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    txBox.text_frame.word_wrap = True
    txBox.text_frame.vertical_anchor = anchor
    p = txBox.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.alignment = align
    return txBox


def add_bullets(slide, left, top, width, height, items, size=16, color=WHITE):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.space_after = Pt(8)
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.text = item
    return txBox


def add_image(slide, path, left, top, width=None, height=None):
    kwargs = {}
    if width:
        kwargs["width"] = Inches(width)
    if height:
        kwargs["height"] = Inches(height)
    return slide.shapes.add_picture(str(path), Inches(left), Inches(top), **kwargs)


def add_card(slide, left, top, width, height):
    shape = slide.shapes.add_shape(
        1, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = DARK_CARD
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


# ── Slide 1: Title ──
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)
add_text(s, 1, 1.8, 11, 1.2, "MealPlanAgent", size=48, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)
add_text(s, 1, 3.2, 11, 0.8, "Hybrid RAG with Knowledge Graph Re-ranking\nfor Recipe Recommendation",
         size=24, color=WHITE, align=PP_ALIGN.CENTER)
add_text(s, 1, 5.0, 11, 0.6, "CMPE 258  |  Deep Learning  |  San Jose State University",
         size=16, color=LIGHT_GRAY, align=PP_ALIGN.CENTER)

# ── Slide 2: Problem ──
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)
add_text(s, 0.8, 0.4, 6, 0.7, "The Problem", size=32, color=ACCENT, bold=True)
add_card(s, 0.8, 1.3, 5.5, 2.5)
add_text(s, 1.1, 1.5, 5, 2.2,
         "LLMs hallucinate recipes:\n"
         "  Wrong cooking times\n"
         "  Impossible ingredient combos\n"
         "  Fabricated nutrition data\n"
         "  \"Nut-free\" with almond flour",
         size=18, color=WHITE)
add_card(s, 0.8, 4.1, 5.5, 2.5)
add_text(s, 1.1, 4.3, 5, 2.2,
         "Our solution: RAG\n"
         "  Ground every recipe in a real database\n"
         "  10,000 Food.com recipes\n"
         "  Verified nutrition, traceable citations\n"
         "  LLM reasons, code verifies",
         size=18, color=WHITE)
add_image(s, CHARTS / "knowledge_graph_single_recipe.png", 7.0, 1.0, width=5.8)

# ── Slide 3: Architecture ──
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)
add_text(s, 0.8, 0.4, 11, 0.7, "System Architecture", size=32, color=ACCENT, bold=True)
add_card(s, 0.8, 1.3, 11.7, 5.5)
add_text(s, 1.2, 1.5, 11, 5.2,
         "User Input (NL or structured)\n"
         "        |\n"
         "   NL Parser  -->  extracts: ingredients, cuisine, allergens, time, calories\n"
         "        |\n"
         "   Planner (LLM)  -->  generates meal queries with preferred_ingredients/tags\n"
         "        |\n"
         "   Executor  -->  5 deterministic tools\n"
         "        |           recipe_search (BM25 + Vector + KG)\n"
         "        |           allergy_checker, nutrition, grocery, budget\n"
         "        |\n"
         "   Critic  -->  rule-based verification (5 checks)\n"
         "        |\n"
         "   If invalid: retry (max 2)  |  If valid: return results",
         size=16, color=WHITE)

# ── Slide 4: Hybrid RAG Pipeline ──
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)
add_text(s, 0.8, 0.4, 11, 0.7, "Three-Stage Hybrid Retrieval", size=32, color=ACCENT, bold=True)

# BM25 card
add_card(s, 0.5, 1.3, 3.8, 3.0)
add_text(s, 0.7, 1.4, 3.5, 0.5, "1. BM25 (Sparse)", size=20, color=ACCENT, bold=True)
add_text(s, 0.7, 1.9, 3.5, 2.2,
         "Exact term matching\n"
         "\"vegetarian\" matches tag\n"
         "IDF weighting: rare terms\n"
         "score more\n"
         "P@5 = 0.575 alone",
         size=15, color=WHITE)

# Vector card
add_card(s, 4.7, 1.3, 3.8, 3.0)
add_text(s, 4.9, 1.4, 3.5, 0.5, "2. Vector (Dense)", size=20, color=ACCENT, bold=True)
add_text(s, 4.9, 1.9, 3.5, 2.2,
         "Semantic similarity\n"
         "all-MiniLM-L6 (384d, 22M)\n"
         "ChromaDB + HNSW index\n"
         "Captures paraphrases\n"
         "P@5 = 0.255 alone",
         size=15, color=WHITE)

# RRF + KG card
add_card(s, 8.9, 1.3, 3.8, 3.0)
add_text(s, 9.1, 1.4, 3.5, 0.5, "3. RRF + KG Re-rank", size=20, color=ACCENT, bold=True)
add_text(s, 9.1, 1.9, 3.5, 2.2,
         "RRF fuses both ranked lists\n"
         "score = 1/(k + rank + 1)\n"
         "KG boosts by ingredient\n"
         "overlap with preferences\n"
         "Full hybrid P@5 = 0.655",
         size=15, color=WHITE)

# Arrow + result
add_text(s, 0.5, 4.5, 12, 0.5,
         "BM25 candidates  +  Vector candidates  -->  RRF fusion  -->  KG re-ranking  -->  Top-k results",
         size=16, color=LIGHT_GRAY, align=PP_ALIGN.CENTER)

add_image(s, CHARTS / "component_contribution.png", 1.5, 5.0, width=10)

# ── Slide 5: Knowledge Graph ──
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)
add_text(s, 0.8, 0.4, 6, 0.7, "Knowledge Graph", size=32, color=ACCENT, bold=True)
add_text(s, 0.8, 1.0, 5, 0.4, "15,807 nodes  |  270,466 edges  |  3 node types", size=16, color=LIGHT_GRAY)
add_image(s, CHARTS / "kg_analysis.png", 0.3, 1.5, width=12.5)

# ── Slide 6: Ablation Study ──
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)
add_text(s, 0.8, 0.4, 11, 0.7, "Ablation Study Results", size=32, color=ACCENT, bold=True)
add_text(s, 0.8, 0.9, 6, 0.4, "40 stratified queries across 5 types", size=14, color=LIGHT_GRAY)
add_image(s, CHARTS / "ablation_comparison.png", 0.3, 1.3, width=12.5)

# ── Slide 7: Statistical Confidence ──
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)
add_text(s, 0.8, 0.4, 11, 0.7, "Statistical Significance", size=32, color=ACCENT, bold=True)
add_image(s, CHARTS / "ablation_with_ci.png", 0.3, 1.2, width=12.5)
add_card(s, 0.8, 5.8, 11.7, 1.3)
add_text(s, 1.1, 5.9, 11, 1.1,
         "full_hybrid vs vector_only (P@5):  d = 1.225 (large),  p < 0.001\n"
         "KG disabled (w=0) drops P@5 by 44% and NDCG@10 by 47%\n"
         "4 of 5 metrics pass Bonferroni correction (alpha = 0.0033)  |  n = 40",
         size=14, color=WHITE)

# ── Slide 8: Hyperparameter Sensitivity ──
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)
add_text(s, 0.8, 0.4, 11, 0.7, "Hyperparameter Sensitivity", size=32, color=ACCENT, bold=True)
add_image(s, CHARTS / "hyperparam_sensitivity.png", 0.3, 1.2, width=12.5)
add_card(s, 0.8, 5.5, 11.7, 1.5)
add_bullets(s, 1.1, 5.6, 11, 1.3, [
    "top_k: most impactful. k=5 to k=50 improves NDCG@10 by 156%",
    "RRF k: robust across [30, 200]. Default k=60 is optimal",
    "KG weight: binary effect. Any w > 0 restores full performance; no degradation at w=1.0",
], size=14, color=WHITE)

# ── Slide 9: Query Type Breakdown ──
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)
add_text(s, 0.8, 0.4, 11, 0.7, "Performance by Query Type", size=32, color=ACCENT, bold=True)
add_image(s, CHARTS / "query_type_breakdown.png", 0.3, 1.2, width=7.5)
add_card(s, 8.2, 1.2, 4.5, 5.5)
add_bullets(s, 8.5, 1.5, 4, 5, [
    "Tag-single: all configs\ndo well (MRR 0.87-1.0)",
    "",
    "Ingredient queries: KG\nmakes the biggest difference\n(MRR 0.17 to 0.75)",
    "",
    "NL queries: BM25-based\nconfigs lead (MRR 0.56)\nFull hybrid drops to 0.33",
    "",
    "Full hybrid wins aggregate\nbut has NL weakness;\nKG is what makes fusion work",
], size=14, color=WHITE)

# ── Slide 10: Embedding Space ──
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)
add_text(s, 0.8, 0.4, 11, 0.7, "Embedding Space Visualization", size=32, color=ACCENT, bold=True)
add_image(s, CHARTS / "embedding_space_umap.png", 0.2, 1.0, height=5.8)
add_image(s, CHARTS / "embedding_space_tsne.png", 6.6, 1.0, height=5.8)
add_text(s, 0.2, 6.9, 6, 0.4, "UMAP  |  Silhouette: -0.269", size=14, color=LIGHT_GRAY, align=PP_ALIGN.CENTER)
add_text(s, 6.6, 6.9, 6, 0.4, "t-SNE  |  2,000 recipes, 11 cuisines", size=14, color=LIGHT_GRAY, align=PP_ALIGN.CENTER)

# ── Slide 11: P-R Curves + Latency ──
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)
add_text(s, 0.8, 0.4, 11, 0.7, "Precision-Recall & Latency", size=32, color=ACCENT, bold=True)
add_image(s, CHARTS / "precision_recall_curves.png", 0.2, 1.1, width=6.3)
add_image(s, CHARTS / "retrieval_latency.png", 6.8, 1.1, width=6.0)

# ── Slide 12: Conclusions ──
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)
add_text(s, 0.8, 0.4, 11, 0.7, "Key Takeaways", size=32, color=ACCENT, bold=True)

conclusions = [
    ("Full hybrid wins every aggregate metric",
     "P@5=0.655, MRR=0.688, NDCG@10=0.438\n+14% over BM25-only; d=1.225 on P@5 (p<0.001)"),
    ("KG makes fusion work; without it, RRF hurts",
     "RRF without KG drops P@5 36% below BM25-only\nKG re-ranking recovers and surpasses (+14%)"),
    ("top_k is the most impactful hyperparameter",
     "k=5 to k=50: NDCG@10 +156%, AP +189%\nDeeper retrieval amplifies KG effectiveness"),
    ("Each component has genuine strengths and weaknesses",
     "BM25 leads NL queries (0.56 vs 0.33); KG drives ingredients\nFull hybrid wins aggregate, not every query type"),
]

for i, (title, detail) in enumerate(conclusions):
    y = 1.2 + i * 1.5
    add_card(s, 0.8, y, 11.7, 1.3)
    add_text(s, 1.1, y + 0.1, 11, 0.4, title, size=20, color=ACCENT, bold=True)
    add_text(s, 1.1, y + 0.55, 11, 0.7, detail, size=14, color=LIGHT_GRAY)


out = Path("report/Presentation.pptx")
prs.save(out)
print(f"Saved {out} ({len(prs.slides)} slides)")
